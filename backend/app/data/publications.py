"""Deterministic adapters for inspected DOSM workbooks; no transcribed values."""
import re
from io import BytesIO
import polars as pl
from openpyxl import load_workbook
from .states import normalize_state, STATES, TERRITORIES

def parse_workbook(content, adapter):
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    if adapter == 'gdp_capita':
        sheet = workbook['A22-A23']
        rows = list(sheet.iter_rows(values_only=True))
        start = next(i for i,r in enumerate(rows) if any('GDP per capita by state at current prices' in str(v) for v in r))
        header = start + 2
        field = 'gdp_per_capita'
    elif adapter == 'productivity':
        sheet = workbook['Jadual 1']; rows = list(sheet.iter_rows(values_only=True)); header = 3
        if 'value added per employment' not in str(rows[1]):
            raise ValueError('Productivity workbook schema changed')
        field = 'labour_productivity'
    else:
        raise ValueError('Unknown workbook adapter')
    years = {i: int(re.match(r'^(\d{4})',str(v))[1]) for i,v in enumerate(rows[header]) if re.match(r'^\d{4}',str(v))}
    if not years: raise ValueError('Missing workbook year header')
    result=[]
    for row_index,row in enumerate(rows[header+1:],start=header+2):
        state=normalize_state(str(row[1] or ''))
        if state not in STATES + TERRITORIES + ['Malaysia']: continue
        for col,year in years.items():
            value=row[col]
            if value is not None:
                if not isinstance(value,(int,float)): raise ValueError('Non-numeric publication cell')
                result.append({'state':state,'date':f'{year}-01-01',field:float(value),'sheet':sheet.title,'cell':f'{sheet.cell(row_index,col+1).column_letter}{row_index}','year_label':str(rows[header][col])})
    workbook.close()
    if not result: raise ValueError('No publication observations')
    return pl.DataFrame(result)
