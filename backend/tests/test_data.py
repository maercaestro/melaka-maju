from io import BytesIO
import httpx
import polars as pl
import pytest
from openpyxl import Workbook
from app.data import opendosm
from app.data.publications import parse_workbook

def test_cache_no_network_and_metadata(tmp_path,monkeypatch):
    monkeypatch.setenv('DATA_CACHE_DIR',str(tmp_path))
    df=pl.DataFrame({'state':['Melaka'],'date':['2024-01-01'],'gini':[.3]})
    data=BytesIO();df.write_parquet(data)
    calls=[]
    def get(self,url,**kwargs):
        calls.append(url)
        return httpx.Response(200,content=data.getvalue(),request=httpx.Request('GET',url))
    monkeypatch.setattr(httpx.Client,'get',get)
    opendosm.load_dataset('hh_inequality_state')
    opendosm.load_dataset('hh_inequality_state')
    assert len(calls)==1
    metadata=opendosm.get_dataset_metadata('hh_inequality_state')
    assert metadata['max_year']==2024 and metadata['sha256']
    assert metadata['retrieved_at']

def test_failed_refresh_keeps_good_cache(tmp_path,monkeypatch):
    monkeypatch.setenv('DATA_CACHE_DIR',str(tmp_path))
    path=tmp_path/'hh_inequality_state.parquet'
    pl.DataFrame({'state':['Melaka'],'date':['2024-01-01'],'gini':[.3]}).write_parquet(path)
    original=path.read_bytes()
    def fail(*args,**kwargs):raise httpx.ConnectError('offline')
    monkeypatch.setattr(httpx.Client,'get',fail)
    with pytest.raises(opendosm.DataUnavailable):opendosm.refresh_dataset('hh_inequality_state')
    assert path.read_bytes()==original

def test_source_validation():
    for row in [{'state':'Atlantis','date':'2024-01-01'},{'state':None,'date':'2024-01-01'},{'state':'Melaka','date':None}]:
        with pytest.raises(ValueError):opendosm.validate_source(pl.DataFrame([row]))

def test_workbook_years_and_coordinates():
    w=Workbook();s=w.active;s.title='Jadual 1'
    s.cell(2,2,'Table 1: value added per employment');s.cell(4,3,2022);s.cell(4,4,2024)
    s.cell(5,2,'Melaka');s.cell(5,3,100);s.cell(5,4,120)
    stream=BytesIO();w.save(stream)
    df=parse_workbook(stream.getvalue(),'productivity')
    assert df['date'].to_list()==['2022-01-01','2024-01-01']
    assert df['cell'].to_list()==['C5','D5']

def test_conflicting_source_dimensions_rejected():
    df=pl.DataFrame({"state":["Melaka","Melaka"],"date":["2024-01-01"]*2,"gini":[.3,.4]})
    with pytest.raises(ValueError,match="Duplicate source dimension"):
        opendosm.validate_source(df)
